"""Atualiza o painel de evidências da aba Ativos a partir de fotos/TAG e checklists/.

Uso: py -3 atualizar_evidencias.py [--planilha Controle_Ativos.xlsx]
"""

from __future__ import annotations

import argparse
import os
import re
import tempfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parent
EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp", ".tif", ".tiff"}
HEADERS = (
    "2.1\nfoto_ativo", "7.1\nfoto_antes", "9.2\nfoto_painel",
    "9.3\nfoto_gnss", "9.4\nfoto_chicotes1",
    "9.4\nfoto_chicotes2 (opcional)", "9.4\nfoto_chicotes3 (opcional)",
    "9.4\nfoto_chicotes4 (opcional)", "10.1\nprint_classe",
    "10.2\nprint_configuracao", "10.3\nfoto_comissionado",
    "11.1\nchecklist (PDF/imagem)",
    *(f"7.2\nfoto_antes_adicional{i} (opcional)" for i in range(1, 9)),
)
GREEN = PatternFill("solid", fgColor="C6EFCE")
RED = PatternFill("solid", fgColor="FFC7CE")
NONE = PatternFill(fill_type=None)
HEADER = PatternFill("solid", fgColor="1F4E78")
OK_FONT = Font(name="Calibri", size=13, bold=True, color="006100")
NO_FONT = Font(name="Calibri", size=13, bold=True, color="9C0006")
NEUTRAL_FONT = Font(name="Calibri", size=11, color="505A64")
HEADER_FONT = Font(name="Calibri", size=10, bold=True, color="FFFFFF")


def image_files(folder: Path, stem: str) -> list[Path]:
    if not folder.is_dir():
        return []
    return [p for p in folder.iterdir() if p.is_file()
            and p.stem.casefold() == stem.casefold()
            and p.suffix.casefold() in EXTENSIONS]


def checklist_pdfs(tag: str, folder: Path, checklists: Path) -> list[Path]:
    prefix = re.compile(r"^" + re.escape(tag) + r"(?:\s*[-_–]\s*|\s+)checklist\b", re.I)
    found: list[Path] = []
    for directory in (checklists, folder):
        if directory.is_dir():
            found.extend(p for p in directory.iterdir() if p.is_file()
                         and p.suffix.casefold() == ".pdf" and prefix.search(p.stem))
    return found


def chicotes_status(folder: Path) -> tuple[list[str], list[str]]:
    slots = [image_files(folder, f"foto_chicotes{i}") for i in range(1, 5)]
    old = image_files(folder, "foto_chicotes")
    issues = []
    if old and any(slots):
        issues.append("foto_chicotes: padrão antigo misturado com fotos numeradas")
        return ["!" if files or (i == 0 and old) else "" for i, files in enumerate(slots)], issues
    if len(old) > 1:
        issues.append("foto_chicotes: mais de uma imagem com o nome antigo")
        return ["!", "", "", ""], issues
    if old:
        return ["✓", "", "", ""], issues
    for i, files in enumerate(slots, 1):
        if len(files) > 1:
            issues.append(f"foto_chicotes{i}: mais de uma extensão/arquivo")
    return ["!" if len(files) > 1 else "✓" if files else "" for files in slots], issues


def fotos_antes_status(folder: Path) -> tuple[list[str], list[str]]:
    slots = [image_files(folder, f"foto_antes_adicional{i}") for i in range(1, 9)]
    issues = []
    if folder.is_dir():
        for path in folder.iterdir():
            if path.is_file() and path.stem.casefold().startswith("foto_antes_adicional"):
                suffix = path.stem.casefold().removeprefix("foto_antes_adicional")
                if suffix not in tuple(str(i) for i in range(1, 9)) or path.suffix.casefold() not in EXTENSIONS:
                    issues.append(f"foto_antes_adicional: nome/extensão inválidos em {path.name}")
    states = ["!" if len(files) > 1 else "✓" if files else "" for files in slots]
    for i, files in enumerate(slots, 1):
        if len(files) > 1:
            issues.append(f"foto_antes_adicional{i}: mais de um arquivo")
    if issues and not any(state == "!" for state in states):
        states[0] = "!"
    return states, issues


def status(tag: str, folder: Path, checklists: Path) -> tuple[list[str], list[str]]:
    issues: list[str] = []

    def single(stem: str) -> str:
        files = image_files(folder, stem)
        if len(files) > 1:
            issues.append(f"{stem}: mais de uma extensão/arquivo")
            return "!"
        return "✓" if files else ""

    pre = [single(stem) for stem in ("foto_ativo", "foto_antes", "foto_painel", "foto_gnss")]
    chicotes, chicotes_issues = chicotes_status(folder)
    issues.extend(chicotes_issues)
    post = [single(stem) for stem in ("print_classe", "print_configuracao", "foto_comissionado")]
    pdfs = checklist_pdfs(tag, folder, checklists)
    if len(pdfs) > 1:
        issues.append("checklist: mais de um PDF correspondente à TAG")
        checklist = "!"
    elif pdfs:
        checklist = "✓"  # PDF tem prioridade sobre a imagem no gerador.
    else:
        checklist = single("checklist_assinado")
    additional, additional_issues = fotos_antes_status(folder)
    issues.extend(additional_issues)
    return pre + chicotes + post + [checklist] + additional, issues


def update(planilha: Path, fotos: Path, checklists: Path) -> tuple[int, int, list[str]]:
    wb = load_workbook(planilha)
    if "Ativos" not in wb:
        raise ValueError("A planilha precisa conter a aba 'Ativos'.")
    ws = wb["Ativos"]
    if (ws["B1"].value or "").strip().upper() != "TAG":
        raise ValueError("A coluna B da aba 'Ativos' precisa conter TAG no cabeçalho.")

    for j, label in enumerate(HEADERS, 12):
        col = get_column_letter(j)
        cell = ws.cell(1, j, label)
        cell.fill = HEADER
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[col].width = 27 if j >= 24 else 23 if j in (17, 18, 19, 23) else 19
    ws.column_dimensions["K"].width = 3
    ws.row_dimensions[1].height = max(ws.row_dimensions[1].height or 15, 42)
    ws.freeze_panes = "C2"
    warnings: list[str] = []
    processed = ready = 0
    last_row = 1
    for row in range(2, ws.max_row + 1):
        tag = str(ws.cell(row, 2).value or "").strip()
        if not tag:
            continue
        processed += 1
        last_row = row
        states, issues = status(tag, fotos / tag, checklists)
        for pos, state in enumerate(states):
            col = 12 + pos
            cell = ws.cell(row, col)
            optional = pos in (5, 6, 7) or pos >= 12
            visible = state if state else "" if optional else "✗"
            cell.value = visible
            cell.fill = GREEN if state == "✓" else NONE if optional and not state else RED
            cell.font = OK_FONT if state == "✓" else NEUTRAL_FONT if optional and not state else NO_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ready += all(states[i] == "✓" for i in (0, 1, 2, 3, 4, 8, 9, 10, 11))
        warnings.extend(f"{tag}: {issue}" for issue in issues)
    ws.auto_filter.ref = f"A1:AE{last_row}"
    if "Instrucoes" in wb:
        instructions = wb["Instrucoes"]
        old = instructions["A5"].value
        if isinstance(old, str) and old.startswith("3. Coloque as nove evidências"):
            instructions["A5"] = (
                "3. Coloque as evidências obrigatórias em fotos/TAG/. A Figura 7.2 aceita "
                "até oito fotos opcionais, foto_antes_adicional1 a foto_antes_adicional8. "
                "Veja INSTRUCOES.md e atualize os checks com ATUALIZAR_EVIDENCIAS.bat."
            )

    # Salva primeiro um arquivo temporário; se o Excel estiver com a planilha aberta,
    # a troca falha sem destruir o arquivo original.
    fd, temporary = tempfile.mkstemp(prefix="evidencias_", suffix=".xlsx", dir=planilha.parent)
    os.close(fd)
    try:
        wb.save(temporary)
        os.replace(temporary, planilha)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)
    return processed, ready, warnings


def main() -> int:
    parser = argparse.ArgumentParser(description="Atualiza os checks das evidências na aba Ativos.")
    parser.add_argument("--planilha", type=Path, default=ROOT / "Controle_Ativos.xlsx")
    parser.add_argument("--fotos", type=Path, default=ROOT / "fotos")
    parser.add_argument("--checklists", type=Path, default=ROOT / "checklists")
    args = parser.parse_args()
    try:
        count, ready, issues = update(args.planilha, args.fotos, args.checklists)
    except (OSError, ValueError) as exc:
        print(f"ERRO: {exc}\nFeche o Excel, confira os caminhos e tente novamente.")
        return 1
    print(f"Atualizado: {args.planilha} | TAGs: {count} | evidências obrigatórias completas: {ready}")
    for issue in issues:
        print(f"ATENÇÃO: {issue}")
    print("✓ verde = encontrada | ✗ vermelho = falta | ! vermelho = arquivos conflitantes")
    print("Chicotes 2 a 4 e fotos adicionais 7.2 vazios ficam sem cor; checklist aceita imagem ou PDF por TAG.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
